# -*- coding: utf-8 -*-
"""Ara Rapor 3 (Markdown) -> .docx

Turk akademik bicimi: A4, Times New Roman 12 pt, sol kenar 3 cm, digerleri 2,5 cm,
1,15 satir araligi, iki yana yasli govde, otomatik icindekiler alani, sayfa numarasi.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GIRDI, CIKTI = sys.argv[1], sys.argv[2]
GIRDI_KLASOR = Path(GIRDI).resolve().parent     # görsel yolları buna göre çözülür
SATIRLAR = open(GIRDI, encoding="utf-8").read().replace("\r\n", "\n").split("\n")

# Sayfa yüksekliği 29,7 − 2×2,5 = 24,7 cm. Şekil bundan büyükse tek başına sayfaya sığmaz;
# üst/alt boşluk ve başlıkla birlikte sığması için 21 cm ile sınırlıyoruz.
AZAMI_SEKIL_CM = 21.0

doc = Document()

# ----------------------------------------------------------------- sayfa ve stil
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.0), Cm(29.7)
s.left_margin, s.right_margin = Cm(3.0), Cm(2.5)
s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.5)
ICERIK_CM = 21.0 - 3.0 - 2.5          # 15,5 cm kullanilabilir genislik

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)

for ad, boyut in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
    st = doc.styles[ad]
    st.font.name = "Times New Roman"
    st.font.size = Pt(boyut)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    st.paragraph_format.keep_with_next = True


def alan(p, kod: str):
    """Word alan kodu ekler (icindekiler, sayfa numarasi gibi)."""
    r = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = kod
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Güncellemek için: alanı seçip F9"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, t, f3):
        r._r.append(e)


def golgele(hucre, renk: str):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")          # SOLID siyah basar, clear dogru olan
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), renk)
    hucre._tc.get_or_add_tcPr().append(sh)


def alt_kenar(p, renk="AAB6C0", kalinlik=6):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(kalinlik))
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), renk)
    bd.append(b); pPr.append(bd)


def sol_kenar(p, renk="7A8B99", kalinlik=12):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:left")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(kalinlik))
    b.set(qn("w:space"), "8"); b.set(qn("w:color"), renk)
    bd.append(b); pPr.append(bd)


# ----------------------------------------------------------------- satir ici bicim
DP = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\n]+\*)")


def yaz_bicimli(p, metin: str, boyut: Pt | None = None):
    son = 0
    for m in DP.finditer(metin):
        if m.start() > son:
            r = p.add_run(metin[son:m.start()])
            if boyut: r.font.size = boyut
        t = m.group(0)
        if t.startswith("**"):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("`"):
            r = p.add_run(t[1:-1]); r.font.name = "Consolas"
            r.font.size = Pt(10) if not boyut else boyut
        else:
            r = p.add_run(t[1:-1]); r.italic = True
        if boyut and not t.startswith("`"): r.font.size = boyut
        son = m.end()
    if son < len(metin):
        r = p.add_run(metin[son:])
        if boyut: r.font.size = boyut


def sade_baslik(metin: str) -> str:
    """Baslik metninden Markdown isaretlerini atar.

    Baslikta kalan ters tirnak Word'de hata gibi gorunuyor ve otomatik icindekiler
    alanina da oyle giriyor; bicimi basliklarda tasimiyoruz, metni sadelestiriyoruz.
    """
    return metin.replace("`", "").replace("**", "").strip()


def hucreler(satir: str) -> list[str]:
    return [h.strip() for h in satir.strip().strip("|").split("|")]


def tablo_yap(blok: list[str]):
    basliklar = hucreler(blok[0])
    govde = [hucreler(x) for x in blok[2:]]
    n = len(basliklar)

    t = doc.add_table(rows=1, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    # Ilk sutun isim tasiyor, digerleri sayi: ilk sutuna daha fazla yer.
    # Genis tabloda ilk sutunun payi dusuyor ve yazi kuculuyor: 9 sutunlu sonuc
    # tablosunda 1,4 cm'lik sutunlara 9,5 pt sigmiyor, basliklar ucer satira bolunuyor.
    if n <= 2:
        gen = [ICERIK_CM / n] * n
        punto = Pt(9.5)
    else:
        pay = 0.20 if n >= 8 else (0.26 if n >= 6 else 0.34)
        ilk = ICERIK_CM * pay
        gen = [ilk] + [(ICERIK_CM - ilk) / (n - 1)] * (n - 1)
        punto = Pt(8.0) if n >= 8 else Pt(9.5)

    def doldur(satir_h, degerler, basliksa):
        for i, h in enumerate(satir_h.cells):
            h.width = Cm(gen[i])
            p = h.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            v = degerler[i] if i < len(degerler) else ""
            if basliksa:
                r = p.add_run(v.replace("**", "").replace("`", "")); r.bold = True; r.font.size = punto
                golgele(h, "E8EDF2")
            else:
                yaz_bicimli(p, v, boyut=punto)

    doldur(t.rows[0], basliklar, True)
    for satir in govde:
        doldur(t.add_row(), satir, False)
    return t


# ----------------------------------------------------------------- cozumleme
i = 0
kapak_bitti = False

while i < len(SATIRLAR):
    ham = SATIRLAR[i]
    k = ham.strip()

    if k == "":
        i += 1
        continue

    # yatay cizgi -> paragraf alt kenari (tablo kullanmiyoruz)
    if re.fullmatch(r"-{3,}", k):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        alt_kenar(p)
        i += 1
        continue

    # kod blogu
    if k.startswith("```"):
        i += 1
        kod = []
        while i < len(SATIRLAR) and not SATIRLAR[i].strip().startswith("```"):
            kod.append(SATIRLAR[i]); i += 1
        i += 1
        for idx, satir in enumerate(kod):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(6 if idx == 0 else 0)
            p.paragraph_format.space_after = Pt(7 if idx == len(kod) - 1 else 0)
            r = p.add_run(satir if satir else " ")
            r.font.name = "Consolas"; r.font.size = Pt(9)
        continue

    # görsel:  ![alt](yol)
    m_img = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", k)
    if m_img:
        yol = (GIRDI_KLASOR / m_img.group(2)).resolve()
        if not yol.exists():
            print(f"UYARI görsel bulunamadi, atlandi: {yol}")
            i += 1
            continue
        with Image.open(yol) as im:
            g_px, y_px = im.size
        # Genişlik sayfa genişliğine oturur; ama uzun şekiller sayfa yüksekliğini aşıyor,
        # o yüzden yükseklik sınırını da kontrol edip gerekirse genişliği küçültüyoruz.
        g_cm = ICERIK_CM
        y_cm = y_px / g_px * g_cm
        if y_cm > AZAMI_SEKIL_CM:
            g_cm *= AZAMI_SEKIL_CM / y_cm
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.keep_together = True
        p.add_run().add_picture(str(yol), width=Cm(g_cm))
        i += 1
        continue

    # tablo
    if k.startswith("|") and i + 1 < len(SATIRLAR) and re.fullmatch(r"\|[\s:|-]+\|", SATIRLAR[i + 1].strip()):
        blok = []
        while i < len(SATIRLAR) and SATIRLAR[i].strip().startswith("|"):
            blok.append(SATIRLAR[i]); i += 1
        tablo_yap(blok)
        bos = doc.add_paragraph()
        bos.paragraph_format.space_after = Pt(8)
        continue

    # basliklar
    if k.startswith("# "):
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(sade_baslik(k[2:]))
        r.font.name = "Times New Roman"; r.font.size = Pt(20); r.bold = True
        p.paragraph_format.space_after = Pt(18)
        i += 1
        continue

    if k.startswith("## "):
        metin = k[3:]
        if not kapak_bitti and re.match(r"^1\.", metin):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            ic = doc.add_paragraph(style="Heading 1")
            ic.add_run("İÇİNDEKİLER")
            pi = doc.add_paragraph()
            alan(pi, r'TOC \o "1-2" \h \z \u')
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            kapak_bitti = True
        p = doc.add_paragraph(style="Heading 1")
        p.add_run(sade_baslik(metin))
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(9)
        i += 1
        continue

    if k.startswith("### "):
        p = doc.add_paragraph(style="Heading 2")
        p.add_run(sade_baslik(k[4:]))
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(7)
        i += 1
        continue

    # alinti
    if k.startswith(">"):
        alinti = []
        while i < len(SATIRLAR) and SATIRLAR[i].strip().startswith(">"):
            alinti.append(re.sub(r"^>\s?", "", SATIRLAR[i].strip())); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(8)
        sol_kenar(p)
        yaz_bicimli(p, " ".join(x for x in alinti if x))
        for r in p.runs:
            r.italic = True
        continue

    # madde isaretli liste
    if re.match(r"^[-*] ", k):
        ogeler = []
        while i < len(SATIRLAR):
            t = SATIRLAR[i].strip()
            if re.match(r"^[-*] ", t):
                ogeler.append(t[2:]); i += 1
            elif t and re.match(r"^\s{2,}\S", SATIRLAR[i]) and ogeler:
                ogeler[-1] += " " + t; i += 1
            else:
                break
        for o in ogeler:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            yaz_bicimli(p, o)
        continue

    # numarali liste
    if re.match(r"^\d+\. ", k):
        ogeler = []
        while i < len(SATIRLAR):
            t = SATIRLAR[i].strip()
            if re.match(r"^\d+\. ", t):
                ogeler.append(re.sub(r"^\d+\.\s", "", t)); i += 1
            elif t and re.match(r"^\s{3,}\S", SATIRLAR[i]) and ogeler:
                ogeler[-1] += " " + t; i += 1
            else:
                break
        for o in ogeler:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            yaz_bicimli(p, o)
        continue

    # normal paragraf
    kaba = []
    while i < len(SATIRLAR):
        t = SATIRLAR[i]
        kk = t.strip()
        if (not kk or re.match(r"^#{1,3} ", kk) or kk.startswith("|") or kk.startswith("```")
                or kk.startswith(">") or re.match(r"^[-*] ", kk) or re.match(r"^\d+\. ", kk)
                or re.fullmatch(r"-{3,}", kk)):
            break
        kaba.append(kk); i += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    yaz_bicimli(p, " ".join(kaba))

# ----------------------------------------------------------------- alt bilgi: sayfa no
alt = doc.sections[0].footer.paragraphs[0]
alt.alignment = WD_ALIGN_PARAGRAPH.CENTER
alan(alt, "PAGE")
for r in alt.runs:
    r.font.size = Pt(10)
    r.font.name = "Times New Roman"

# ----------------------------------------------------------------- ozellikler
cp = doc.core_properties
cp.title = "3. Ara Rapor — Uygulama"
cp.author = "Baran Yeşilyurt, Ahmet Görkem Yavuz, Ertuğrul Kabaoğlu"
cp.subject = "BSM 401 Tasarım Projesi"
cp.comments = ("Alzheimer Hastalığının Erken Teşhisi İçin 3B Derin Öğrenme Tabanlı "
               "Karar Destek Sistemi — uygulama raporu")

doc.save(CIKTI)
print(f"OK {CIKTI}")
print(f"   paragraf {len(doc.paragraphs)} · tablo {len(doc.tables)}")
